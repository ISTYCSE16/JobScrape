"""Scored DataFrame -> formatted Excel report + CSV export."""

import os
from datetime import date

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUTPUT_DIR = "output"
SHEET_NAME = "Job Matches"
DESCRIPTION_TRUNCATE = 500

HEADER_FILL = PatternFill(start_color="ADD8E6", end_color="ADD8E6", fill_type="solid")
HEADER_FONT = Font(bold=True)

SCORE_GREEN = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
SCORE_YELLOW = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
SCORE_RED = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

COLUMNS = [
    ("Company", "company", 25),
    ("Job Title", "title", 30),
    ("Location", "location", 20),
    ("Site", "site", 15),
    ("Hours Since Posted", "hours_since_posted", 18),
    ("Match Score", "match_score", 15),
    ("Strong Keywords", "strong_keywords", 30),
    ("Missing Keywords", "missing_keywords", 30),
    ("CV Changes Summary", "cv_changes_summary", 40),
    ("Job URL", "job_url", 20),
    ("Description", "description", 50),
]


def _to_comma_string(value) -> str:
    if isinstance(value, (list, tuple)):
        return ", ".join(str(v) for v in value)
    if pd.isna(value):
        return ""
    return str(value)


def _score_fill(score) -> PatternFill:
    try:
        score = float(score)
    except (TypeError, ValueError):
        return None
    if score >= 70:
        return SCORE_GREEN
    if score >= 40:
        return SCORE_YELLOW
    return SCORE_RED


def generate_report(df: pd.DataFrame, output_dir: str = OUTPUT_DIR) -> str:
    os.makedirs(output_dir, exist_ok=True)
    today = date.today().isoformat()

    csv_path = os.path.join(output_dir, f"jobs_{today}.csv")
    df.to_csv(csv_path, index=False)

    xlsx_path = os.path.join(output_dir, f"jobs_{today}.xlsx")

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME

    headers = [label for label, _, _ in COLUMNS]
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    score_col_idx = next(
        i for i, (_, field, _) in enumerate(COLUMNS, start=1) if field == "match_score"
    )
    url_col_idx = next(
        i for i, (_, field, _) in enumerate(COLUMNS, start=1) if field == "job_url"
    )
    desc_col_idx = next(
        i for i, (_, field, _) in enumerate(COLUMNS, start=1) if field == "description"
    )

    for row_idx, row in enumerate(df.itertuples(index=False), start=2):
        for col_idx, (_, field, _) in enumerate(COLUMNS, start=1):
            value = getattr(row, field, "")

            if field in ("strong_keywords", "missing_keywords"):
                value = _to_comma_string(value)
            elif field == "description":
                value = "" if pd.isna(value) else str(value)[:DESCRIPTION_TRUNCATE]
            elif pd.isna(value) if not isinstance(value, (list, tuple)) else False:
                value = ""

            cell = ws.cell(row=row_idx, column=col_idx, value=value)

            if col_idx == score_col_idx:
                fill = _score_fill(value)
                if fill:
                    cell.fill = fill
            elif col_idx == url_col_idx and value:
                cell.hyperlink = value
                cell.font = Font(color="0563C1", underline="single")
            elif col_idx == desc_col_idx:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    for col_idx, (_, _, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"

    wb.save(xlsx_path)
    print(f"Report saved: {xlsx_path}")
    print(f"CSV saved: {csv_path}")
    return xlsx_path


if __name__ == "__main__":
    sample = pd.DataFrame(
        [
            {
                "company": "Example Corp",
                "title": "Software Developer",
                "location": "Edmonton, AB",
                "site": "linkedin",
                "hours_since_posted": 5,
                "match_score": 85,
                "strong_keywords": ["Python", "CI/CD"],
                "missing_keywords": ["Kubernetes"],
                "cv_changes_summary": "Lead with CI/CD automation experience.",
                "job_url": "https://example.com/job/1",
                "description": "We are looking for a Python developer with ML experience.",
            }
        ]
    )
    generate_report(sample)
